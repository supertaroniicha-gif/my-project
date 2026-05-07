#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
フィードバック管理システム - 配布用テンプレート生成（8列版）
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta

# Excelワークブック作成
wb = Workbook()
wb.remove(wb.active)  # デフォルトシートを削除

# ============ Sheet 1: 使い方 ============
ws_guide = wb.create_sheet('📖 使い方', 0)

guide_content = [
    ['オリエンテーション フィードバック管理システム', ''],
    ['', ''],
    ['このテンプレートについて', ''],
    ['このExcelは、オリエンテーション時のフィードバックを記録し、', ''],
    ['自動で統計分析するためのテンプレートです。', ''],
    ['', ''],
    ['🎯 使い方', ''],
    ['1. 「フィードバック入力」シートにデータを入力してください', ''],
    ['2. Good または More をドロップダウンから選択', ''],
    ['3. Next Action は推奨リストから選択（または自由入力）', ''],
    ['4. Google Apps Script を設定すると自動で統計計算します', ''],
    ['', ''],
    ['📋 列の説明', ''],
    ['日付: フィードバックの日付 (例: 2024-05-04)', ''],
    ['曜日: 自動計算される曜日 (Google Sheet で自動入力)', ''],
    ['項目: フィードバックの分類 (例: コミュニケーション)', ''],
    ['Good/More: 良い点(Good) / 改善点(More)', ''],
    ['詳細: 詳しい説明内容', ''],
    ['Next Action①: 実行すべき1番目のアクション', ''],
    ['Next Action②: 実行すべき2番目のアクション', ''],
    ['FB回数: 何回目のフィードバックか (1, 2, 3...)', ''],
    ['', ''],
    ['🔧 Google Apps Script 設定方法', ''],
    ['1. Google Sheet にこのファイルをアップロード', ''],
    ['2. 拡張機能 → Apps Script を開く', ''],
    ['3. google_apps_script_v3.gs のコードをコピペして実行', ''],
    ['4. 自動で統計ダッシュボードが生成されます', ''],
    ['', ''],
    ['💡 Next Action の推奨リスト（47個）', ''],
    ['アイスブレイク拡張、深掘り強化、質問形式工夫、スキル訓練、', ''],
    ['コミュニケーション強化、共感重視、定量データ活用、ヒアリング深化、', ''],
    ['目標設定共有、笑顔増加、説得フロー改善、権威性活用、', ''],
    ['安心感提供、親和性向上、ネック潰し、ADHD対応、他多数', ''],
]

for row_idx, row_data in enumerate(guide_content, 1):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_guide.cell(row=row_idx, column=col_idx)
        cell.value = value
        if '使い方' in str(value) or 'フィードバック管理' in str(value):
            cell.font = Font(bold=True, size=12)
        elif '📖' in str(value) or '🎯' in str(value) or '📋' in str(value) or '🔧' in str(value) or '💡' in str(value):
            cell.font = Font(bold=True, size=11)

ws_guide.column_dimensions['A'].width = 60
ws_guide.column_dimensions['B'].width = 30

# ============ Sheet 2: フィードバック入力 ============
ws_feedback = wb.create_sheet('フィードバック入力', 1)

headers = ['日付', '曜日', '項目', 'Good/More', '詳細', 'Next Action①', 'Next Action②', 'FB回数']
ws_feedback.append(headers)

# ヘッダーのフォーマット
header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for cell in ws_feedback[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

# サンプルデータ
sample_data = [
    ['2024-05-04', '土', 'コミュニケーション', 'Good', '教育プランナーが父にも子供にも話しかけ・目線を合わせた', 'アイスブレイク拡張', '深掘り強化', '1'],
    ['2024-05-04', '土', '心理的配慮', 'Good', '褒めるだけでなく共感を重視', '共感重視', 'ラポール深化', '1'],
    ['2024-05-04', '土', 'ヒアリング', 'More', '毎回勉強について深掘りが不足', 'ヒアリング深化', '質問形式工夫', '1'],
    ['2024-05-05', '日', '説得フロー', 'More', '1課題にすぐトライの良さをぶつけるのはきつい', '段階的説得', '説得フロー改善', '2'],
]

# Good/More の色
good_fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
more_fill = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')

for row_idx, row_data in enumerate(sample_data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_feedback.cell(row=row_idx, column=col_idx)
        cell.value = value
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        # Good/More 列に色付け（列4）
        if col_idx == 4:
            if value == 'Good':
                cell.fill = good_fill
                cell.font = Font(bold=True, color='155724')
            else:
                cell.fill = more_fill
                cell.font = Font(bold=True, color='856404')
            cell.alignment = Alignment(horizontal='center', vertical='center')

    ws_feedback.row_dimensions[row_idx].height = 30

# 空行を追加（ユーザー用）
for row in range(6, 20):
    for col in range(1, 9):  # 8列対応
        cell = ws_feedback.cell(row=row, column=col)
        cell.border = thin_border
    ws_feedback.row_dimensions[row].height = 25

# 列幅設定
ws_feedback.column_dimensions['A'].width = 12
ws_feedback.column_dimensions['B'].width = 10
ws_feedback.column_dimensions['C'].width = 20
ws_feedback.column_dimensions['D'].width = 12
ws_feedback.column_dimensions['E'].width = 40
ws_feedback.column_dimensions['F'].width = 20
ws_feedback.column_dimensions['G'].width = 20
ws_feedback.column_dimensions['H'].width = 10

# ============ Sheet 3: 統計ダッシュボード（テンプレート） ============
ws_stats = wb.create_sheet('統計ダッシュボード', 2)

stats_content = [
    ['📊 統計ダッシュボード', 'Google Apps Script 設定後に自動生成'],
    ['', ''],
    ['総フィードバック数', '4'],
    ['Good数', '2'],
    ['More数', '2'],
    ['Good率(%)', '50'],
    ['', ''],
    ['Next Action ランキング', ''],
    ['項目', '頻度'],
]

for row_idx, row_data in enumerate(stats_content, 1):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_stats.cell(row=row_idx, column=col_idx)
        cell.value = value
        if '統計ダッシュボード' in str(value):
            cell.font = Font(bold=True, size=14)
        elif row_idx == 9:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color='F3F3F3', end_color='F3F3F3', fill_type='solid')

ws_stats.column_dimensions['A'].width = 25
ws_stats.column_dimensions['B'].width = 15

# ============ Sheet 4: Google Apps Script コード ============
ws_code = wb.create_sheet('Google Apps Script', 3)

code_info = [
    ['Google Apps Script セットアップ', ''],
    ['', ''],
    ['以下のステップを実行してください:', ''],
    ['', ''],
    ['1. Google Sheet に このファイルをアップロード', ''],
    ['2. メニュー: 拡張機能 → Apps Script', ''],
    ['3. デフォルトコードを削除', ''],
    ['4. google_apps_script_v3.gs のコードをコピペ', ''],
    ['5. 保存 (Ctrl+S)', ''],
    ['6. 「initializeSheet」を実行', ''],
    ['7. 権限を許可', ''],
    ['', ''],
    ['📝 コード取得方法:', ''],
    ['GitHub リポジトリから google_apps_script_v3.gs を取得', ''],
    ['または、プロジェクトファイルを確認してください', ''],
    ['', ''],
    ['✅ セットアップ完了後:', ''],
    ['- フィードバック入力シートにデータを入力', ''],
    ['- 曜日が自動計算される', ''],
    ['- 統計ダッシュボードが自動で更新される', ''],
    ['- Next Action がドロップダウンで選択可能', ''],
    ['- 頻度が自動で計算される', ''],
]

for row_idx, row_data in enumerate(code_info, 1):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_code.cell(row=row_idx, column=col_idx)
        cell.value = value
        if 'Google Apps Script' in str(value):
            cell.font = Font(bold=True, size=12)
        elif '1.' in str(value) or '2.' in str(value) or row_idx == 17:
            cell.font = Font(bold=True)

ws_code.column_dimensions['A'].width = 50
ws_code.column_dimensions['B'].width = 20

# ファイルを保存
output_file = '/home/user/my-project/フィードバック管理システム_テンプレートv2.xlsx'
wb.save(output_file)

print(f"✅ テンプレートExcelファイル（8列版）が作成されました！")
print(f"📄 ファイル: {output_file}")
print(f"\n📋 シート構成:")
print(f"   1. 📖 使い方 - セットアップガイド")
print(f"   2. フィードバック入力 - データ入力用（8列、曜日自動計算対応）")
print(f"   3. 統計ダッシュボード - 自動計算結果表示用")
print(f"   4. Google Apps Script - コード設定情報")
print(f"\n📊 列構成:")
print(f"   A: 日付")
print(f"   B: 曜日（自動計算）")
print(f"   C: 項目")
print(f"   D: Good/More")
print(f"   E: 詳細")
print(f"   F: Next Action①")
print(f"   G: Next Action②")
print(f"   H: FB回数")
print(f"\n🎯 配布準備完了！")
