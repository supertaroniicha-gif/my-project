#!/usr/bin/env python3
"""
名刺画像を読み込んで、データを抽出し、Excelにまとめ、
AI営業メール下書きをGmailに保存する。
"""

import os
import base64
import json
import datetime
from pathlib import Path
from urllib.request import urlretrieve

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import anthropic
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

# Gmail の読み書き権限
SCOPES = ["https://www.googleapis.com/auth/gmail.modify"]
CREDENTIALS_FILE = "credentials.json"
TOKEN_FILE = "token.json"

GITHUB_REPO = "https://raw.githubusercontent.com/takpfive/claude-code-course-materials-2026-04/main/cards"


def get_gmail_service():
    """Gmail APIクライアントを認証して返す。"""
    creds = None
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, "w") as token:
            token.write(creds.to_json())

    return build("gmail", "v1", credentials=creds)


def download_cards(output_dir: str = "cards_images"):
    """GitHubから名刺画像をダウンロードする。"""
    Path(output_dir).mkdir(exist_ok=True)

    print("GitHubから名刺画像をダウンロード中...")
    for i in range(1, 21):
        filename = f"card_{i:02d}.jpg"
        url = f"{GITHUB_REPO}/{filename}"
        filepath = os.path.join(output_dir, filename)

        if os.path.exists(filepath):
            print(f"  {filename} （既に存在）")
            continue

        try:
            urlretrieve(url, filepath)
            print(f"  {filename} ✓")
        except Exception as e:
            print(f"  {filename} ✗ ({e})")

    return output_dir


def extract_card_data_with_ai(card_image_path: str) -> dict:
    """Claude AIを使って名刺画像からテキストを抽出する。"""
    client = anthropic.Anthropic()

    with open(card_image_path, "rb") as f:
        image_data = base64.standard_b64encode(f.read()).decode("utf-8")

    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=500,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/jpeg",
                            "data": image_data,
                        },
                    },
                    {
                        "type": "text",
                        "text": """この名刺から以下の情報を日本語で抽出してください。
出力は必ず以下のJSON形式で返してください（コードブロック不要）：

{
  "name": "名前",
  "title": "役職",
  "company": "会社名",
  "email": "メールアドレス（あれば）",
  "phone": "電話番号（あれば）",
  "industry": "業種の推測（金融、IT、不動産など）"
}

もし情報が読み取れない場合は、そのフィールドに「不明」と記入してください。"""
                    }
                ],
            }
        ],
    )

    try:
        text = message.content[0].text.strip()
        if "```" in text:
            text = text.split("```")[1]
            if text.startswith("json"):
                text = text[4:]
        return json.loads(text)
    except (json.JSONDecodeError, IndexError):
        return {
            "name": "不明",
            "title": "不明",
            "company": "不明",
            "email": "不明",
            "phone": "不明",
            "industry": "不明",
        }


def create_excel_from_cards(cards_data: list[dict], output_file: str = "business_cards.xlsx"):
    """抽出したデータをExcelにまとめる。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名刺データ"

    headers = ["No.", "名前", "役職", "会社名", "メール", "電話", "業種"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, card_data in enumerate(cards_data, 1):
        ws.append([
            i,
            card_data.get("name", ""),
            card_data.get("title", ""),
            card_data.get("company", ""),
            card_data.get("email", ""),
            card_data.get("phone", ""),
            card_data.get("industry", ""),
        ])

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 15

    wb.save(output_file)
    print(f"Excel ファイル作成: {output_file}")
    return output_file


def generate_sales_email(card_data: dict) -> str:
    """Claude AIを使ってAI研修営業メールを生成する。"""
    client = anthropic.Anthropic()

    company = card_data.get("company", "貴社")
    title = card_data.get("title", "様")
    industry = card_data.get("industry", "一般企業")

    prompt = f"""
あなたはAI研修サービスを販売するプロフェッショナルです。
以下の情報を持つ人物に対して、AI研修の営業メール本文を日本語で作成してください。

【相手情報】
- 会社: {company}
- 役職: {title}
- 推測される業種: {industry}

【要件】
1. 相手の業種や役職から推測される課題にフォーカス
2. AI研修が具体的にどのような価値をもたらすかを示す
3. 親切で専門的なトーン
4. 長さ：200～300文字
5. 署名は不要（本文のみ）

【出力形式】
メール本文のみを返してください。"""

    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=500,
        messages=[{"role": "user", "content": prompt}],
    )

    return message.content[0].text.strip()


def create_gmail_draft(gmail_service, to_email: str, subject: str, body: str):
    """Gmailに下書きメールを作成する。"""
    from email.mime.text import MIMEText

    message = MIMEText(body, _charset="utf-8")
    message["to"] = to_email
    message["subject"] = subject

    raw_message = base64.urlsafe_b64encode(message.as_bytes()).decode()

    draft_body = {
        "message": {
            "raw": raw_message,
        }
    }

    try:
        gmail_service.users().drafts().create(userId="me", body=draft_body).execute()
        return True
    except Exception as e:
        print(f"  エラー: {e}")
        return False


def main():
    print("=" * 60)
    print("名刺スキャン→Excel化→AI営業メール生成")
    print("=" * 60)

    # 1. 名刺画像をダウンロード
    cards_dir = download_cards()

    # 2. 名刺データを抽出
    print("\n名刺データを抽出中...")
    cards_data = []
    card_files = sorted([f for f in os.listdir(cards_dir) if f.endswith(".jpg")])

    for i, card_file in enumerate(card_files, 1):
        card_path = os.path.join(cards_dir, card_file)
        print(f"  {i}/20: {card_file} を処理中...")
        try:
            card_data = extract_card_data_with_ai(card_path)
            cards_data.append(card_data)
        except Exception as e:
            print(f"    エラー: {e}")
            cards_data.append({
                "name": "不明",
                "title": "不明",
                "company": "不明",
                "email": "不明",
                "phone": "不明",
                "industry": "不明",
            })

    # 3. Excelに出力
    print("\nExcel ファイルを作成中...")
    excel_file = create_excel_from_cards(cards_data)

    # 4. Gmail下書きを作成
    print("\nGmail下書きを作成中...")
    gmail_service = get_gmail_service()

    success_count = 0
    for i, card_data in enumerate(cards_data, 1):
        name = card_data.get("name", "様")
        email = card_data.get("email", "")
        company = card_data.get("company", "")

        if not email or email == "不明":
            print(f"  {i}/20: {name} ({company}) - メールアドレスがないためスキップ")
            continue

        subject = f"AI人材育成プログラムのご提案 - {company}"
        body = generate_sales_email(card_data)

        if create_gmail_draft(gmail_service, email, subject, body):
            print(f"  {i}/20: {name} ({company}) ✓")
            success_count += 1
        else:
            print(f"  {i}/20: {name} ({company}) ✗")

    print("\n" + "=" * 60)
    print(f"完了！")
    print(f"  - Excel ファイル: {excel_file}")
    print(f"  - Gmail下書き: {success_count}件作成")
    print("=" * 60)


if __name__ == "__main__":
    main()
