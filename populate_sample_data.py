#!/usr/bin/env python3
"""
business_cards.xlsx にサンプルデータを入力する。
"""

import openpyxl

SAMPLE_DATA = [
    ("田中太郎", "営業部長", "株式会社デジタルソリューション", "tanaka@digital-sol.jp", "090-1234-5678", "IT・ソフトウェア"),
    ("佐藤花子", "企画室課長", "ABC製造株式会社", "sato@abc-mfg.jp", "090-2345-6789", "製造業"),
    ("山田次郎", "店長", "イオン渋谷店", "yamada@aeon-shibuya.jp", "090-3456-7890", "小売・流通"),
    ("鈴木美咲", "医師", "東京大学病院", "suzuki@todai-hospital.jp", "090-4567-8901", "医療・ヘルスケア"),
    ("伊藤健太", "財務部長", "大和銀行", "itoh@daiwa-bank.jp", "090-5678-9012", "金融・銀行"),
    ("中野由美", "人事課長", "Sony株式会社", "nakano@sony.jp", "090-6789-0123", "電子機器・IT"),
    ("木村拓也", "営業課長", "日本電信電話", "kimura@ntt.jp", "090-7890-1234", "通信・インフラ"),
    ("渡辺美穂", "教育開発部長", "文部科学省", "watanabe@mext.go.jp", "090-8901-2345", "教育・官公庁"),
    ("橋本龍一", "商品企画課長", "トヨタ自動車", "hashimoto@toyota.jp", "090-9012-3456", "自動車・物流"),
    ("田村光子", "営業本部長", "楽天グループ", "tamura@rakuten.jp", "090-0123-4567", "EC・インターネット"),
    ("松本耕平", "CTO", "スタートアップA社", "matsumoto@startup-a.jp", "090-1111-2222", "スタートアップ・IT"),
    ("野村美咲", "マーケティング部長", "資生堂", "nomura@shiseido.jp", "090-2222-3333", "化粧品・美容"),
    ("高田隆一", "営業部課長", "NHK", "takada@nhk.or.jp", "090-3333-4444", "メディア・放送"),
    ("佐々木由紀", "管理部長", "みずほ銀行", "sasaki@mizuho-bank.jp", "090-4444-5555", "金融・銀行"),
    ("小林陽太", "営業開発部長", "アマゾンジャパン", "kobayashi@amazon.jp", "090-5555-6666", "EC・インターネット"),
    ("加藤祐子", "人事部部長", "三菱電機", "kato@mitsubishi-electric.jp", "090-6666-7777", "電子機器・重工業"),
    ("岡本智樹", "営業課長", "ANA", "okamoto@ana.jp", "090-7777-8888", "運輸・物流"),
    ("西村由梨", "企画部課長", "朝日新聞", "nishimura@asahi.jp", "090-8888-9999", "メディア・新聞"),
    ("青木修一", "営業部長", "オリンパス", "aoki@olympus.jp", "090-9999-0000", "精密機器・カメラ"),
    ("吉田由美", "マーケティング課長", "ファーストリテイリング", "yoshida@fastretailing.jp", "090-0000-1111", "小売・ファッション"),
]


def populate_excel():
    """Excelファイルにサンプルデータを入力。"""
    wb = openpyxl.load_workbook("business_cards.xlsx")
    ws = wb.active

    print("サンプルデータをExcelに入力中...")
    for i, (name, title, company, email, phone, industry) in enumerate(SAMPLE_DATA, start=2):
        ws[f"B{i}"] = name
        ws[f"C{i}"] = title
        ws[f"D{i}"] = company
        ws[f"E{i}"] = email
        ws[f"F{i}"] = phone
        ws[f"G{i}"] = industry

    wb.save("business_cards.xlsx")
    print(f"✓ {len(SAMPLE_DATA)}件のサンプルデータを入力しました")


if __name__ == "__main__":
    populate_excel()
