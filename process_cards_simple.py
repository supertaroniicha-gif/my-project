#!/usr/bin/env python3
"""
名刺画像を読み込んで、OCRで抽出し、Excelにまとめる（簡略版）
Gmail連携はこの後で行う
"""

import os
import base64
import json
from pathlib import Path
from urllib.request import urlretrieve
import anthropic
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

GITHUB_REPO = "https://raw.githubusercontent.com/takpfive/claude-code-course-materials-2026-04/main/cards"


def download_cards(output_dir: str = "cards_images"):
    """GitHubから名刺画像をダウンロードする。"""
    Path(output_dir).mkdir(exist_ok=True)

    print("GitHubから名刺画像をダウンロード中...")
    for i in range(1, 21):
        filename = f"card_{i:02d}.jpg"
        url = f"{GITHUB_REPO}/{filename}"
        filepath = os.path.join(output_dir, filename)

        if os.path.exists(filepath):
            print(f"  {filename} （既に存在）")
            continue

        try:
            urlretrieve(url, filepath)
            print(f"  {filename} ✓")
        except Exception as e:
            print(f"  {filename} ✗ ({e})")

    return output_dir


def extract_card_data_with_ai(card_image_path: str) -> dict:
    """Claude AIを使って名刺画像からテキストを抽出する。"""
    client = anthropic.Anthropic()

    with open(card_image_path, "rb") as f:
        image_data = base64.standard_b64encode(f.read()).decode("utf-8")

    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=500,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/jpeg",
                            "data": image_data,
                        },
                    },
                    {
                        "type": "text",
                        "text": """この名刺から以下の情報を日本語で抽出してください。
出力は必ず以下のJSON形式で返してください（コードブロック不要）：

{
  "name": "名前",
  "title": "役職",
  "company": "会社名",
  "email": "メールアドレス（あれば）",
  "phone": "電話番号（あれば）",
  "industry": "業種の推測（金融、IT、不動産など）"
}

もし情報が読み取れない場合は、そのフィールドに「不明」と記入してください。"""
                    }
                ],
            }
        ],
    )

    try:
        text = message.content[0].text.strip()
        if "```" in text:
            text = text.split("```")[1]
            if text.startswith("json"):
                text = text[4:]
        return json.loads(text)
    except (json.JSONDecodeError, IndexError):
        return {
            "name": "不明",
            "title": "不明",
            "company": "不明",
            "email": "不明",
            "phone": "不明",
            "industry": "不明",
        }


def create_excel_from_cards(cards_data: list[dict], output_file: str = "business_cards.xlsx"):
    """抽出したデータをExcelにまとめる。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名刺データ"

    headers = ["No.", "名前", "役職", "会社名", "メール", "電話", "業種"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, card_data in enumerate(cards_data, 1):
        ws.append([
            i,
            card_data.get("name", ""),
            card_data.get("title", ""),
            card_data.get("company", ""),
            card_data.get("email", ""),
            card_data.get("phone", ""),
            card_data.get("industry", ""),
        ])

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 15

    wb.save(output_file)
    print(f"✓ Excel ファイル作成: {output_file}")
    return output_file, cards_data


def main():
    print("=" * 60)
    print("名刺スキャン→Excel化（Google認証不要）")
    print("=" * 60)

    # 1. 名刺画像をダウンロード
    cards_dir = download_cards()

    # 2. 名刺データを抽出
    print("\n名刺データを抽出中（AI処理）...")
    cards_data = []
    card_files = sorted([f for f in os.listdir(cards_dir) if f.endswith(".jpg")])

    for i, card_file in enumerate(card_files, 1):
        card_path = os.path.join(cards_dir, card_file)
        print(f"  [{i:2d}/20] {card_file}...", end=" ", flush=True)
        try:
            card_data = extract_card_data_with_ai(card_path)
            cards_data.append(card_data)
            print("✓")
        except Exception as e:
            print(f"✗ ({e})")
            cards_data.append({
                "name": "不明",
                "title": "不明",
                "company": "不明",
                "email": "不明",
                "phone": "不明",
                "industry": "不明",
            })

    # 3. Excelに出力
    print("\n")
    excel_file, card_data = create_excel_from_cards(cards_data)

    print("\n" + "=" * 60)
    print("✓ 完了！")
    print(f"  - Excel ファイル: {excel_file}")
    print(f"  - 抽出件数: {len(cards_data)}件")
    print("=" * 60)

    # JSONも保存（次のステップで使用）
    with open("cards_data.json", "w", encoding="utf-8") as f:
        json.dump(cards_data, f, ensure_ascii=False, indent=2)
    print(f"✓ JSON保存: cards_data.json")


if __name__ == "__main__":
    main()
